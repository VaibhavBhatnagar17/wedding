#!/usr/bin/env python3
"""Build the Udaipur resort workbook from venues/verified_resorts.py.

    python3.11 venues/build_resort_sheet.py

Only properties that publish at least 40 rooms and both a garden and a banquet
hall reach the shortlist. Availability for 31 Jan – 2 Feb 2027 and the wedding
price are left blank for the phone call.
"""

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from verified_resorts import BAND_RULE, BANDS, CONFIRM_FIRST, NAME_ONLY, REJECTED, SHORTLIST  # noqa: E402

OUT = Path(__file__).resolve().parent / "udaipur-resorts-1-2-feb-2027.xlsx"

MAROON = "47101C"
INK = "23201A"
MUTED = "7B7062"
IVORY = "FDFAF1"
CREAM = "F9F2E2"
YELLOW = "FFF3C4"
GREEN = "E5F2E9"
AMBER = "FDE8C8"
RED = "F8E6E6"
BAND_FILL = {"Affordable": "D8EBD0", "Mid-range": "E2E8F3", "Premium": "EADCF0"}

THIN = Side(style="thin", color="E0D3B8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
BODY = Font(name="Calibri", size=11, color=INK)
BOLD = Font(name="Calibri", size=11, color=INK, bold=True)
HEAD = Font(name="Calibri", size=11, color="F7ECC9", bold=True)
TITLE = Font(name="Calibri", size=18, color=MAROON, bold=True)
NOTE = Font(name="Calibri", size=12, color=MUTED, italic=True)


def fill(hex_):
    return PatternFill("solid", fgColor=hex_)


# (header, key in SHORTLIST, width). None for a yellow input column.
COLUMNS = [
    ("Band", "band", 12),
    ("# in band", "band_rank", 7),
    ("Resort", "name", 34),
    ("Area", "area", 24),
    ("Rooms", "rooms", 11),
    ("Where the room count comes from", "rooms_src", 36),
    ("Garden / lawn, as published", "garden", 42),
    ("Banquet hall, as published", "hall", 38),
    ("Seats 650 for the reception?", "seats_650", 34),
    ("Holds 180 for the core functions?", "seats_180", 28),
    ("Rooms for four adults", "four_adult", 32),
    ("Class, as the hotel describes itself", "star", 26),
    ("Class, as booking sites sell it", "sold_as", 28),
    ("Room rate seen online, 6 Oct 2026", "room_rate", 38),
    ("Rating", "rating", 26),
    ("Plate or package price signal (not a quote)", "price", 36),
    ("Good to know", "notes", 40),
    ("Address", "address", 34),
    ("Phone", "phone", 26),
    ("Email", "email", 30),
    ("Sources", "source", 34),
    ("Confidence", "confidence", 24),
    ("Free 31 Jan 2027", None, 14),
    ("Free 1 Feb 2027", None, 14),
    ("Free 2 Feb 2027", None, 14),
    ("Their room rate, per night", None, 18),
    ("Their quote, total", None, 18),
    ("Notes from the call", None, 32),
]

QUESTIONS = [
    "How many rooms can you contract for the nights of 31 Jan, 1 Feb and 2 Feb 2027? The website total is not a block.",
    "How many rooms have two double beds, or two queens? We need about 15. Name the category and the count.",
    "What is the seated-dinner capacity of the lawn for the 650 reception, after the stage, dance floor and buffets? Floating capacity is not the number.",
    "Which indoor hall seats 180 for dinner, and which seats 650, if it rains?",
    "Are 1 and 2 Feb 2027 free? Will you put a hold in writing?",
    "Will you write 'rooms from 06:00 on 1 Feb 2027, no extra night charged'?",
    "Room rate with breakfast and GST, per night, for two nights. What happens if we do not fill the block, and by when?",
    "Veg plate for a Tuesday, with the minimum number of covers. Is lawn rental separate?",
    "How many complimentary rooms, against what food-and-beverage minimum?",
    "Are outside décor, a DJ, dhol and brass band, and outside alcohol allowed? Any fee?",
    "What time must outdoor music stop, and which indoor room can the after-party move to?",
    "Cancellation and postponement terms, and how much of the advance is forfeited.",
    "Ask for two families who held about 600-guest weddings there, and the lawn they used.",
]


def header_row(ws, headers):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.fill = fill(MAROON)
        cell.font = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    ws.row_dimensions[1].height = 34
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def fit_colour(text):
    t = str(text)
    if t.startswith("Yes"):
        return GREEN
    if t.startswith("No") or t.startswith("Short") or t.startswith("Close"):
        return RED
    return AMBER


def build_start(wb):
    ws = wb.active
    ws.title = "Start here"
    ws.sheet_properties.tabColor = MAROON
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 120
    ws["A1"] = "Udaipur resorts for 1 and 2 February 2027"
    ws["A1"].font = TITLE
    ws["A2"] = "Checked 6 October 2026, in three sweeps. Every row was checked against the rules below."
    ws["A2"].font = NOTE

    counts = {band: sum(1 for s in SHORTLIST if s["band"] == band) for band, _ in BANDS}
    yes = {band: sum(1 for s in SHORTLIST if s["band"] == band and str(s["seats_650"]).startswith("Yes"))
           for band, _ in BANDS}
    checked = len(SHORTLIST) + len(CONFIRM_FIRST) + len(REJECTED) + len(NAME_ONLY)
    lines = [
        ("The rules every shortlisted property passes", True),
        ("At least 40 rooms, published. A garden or lawn and a banquet hall, both published. In or around Udaipur. "
         "Being cheap does not get a property onto the list: affordable resorts with fewer than 40 rooms, or without "
         "a hall, are on 'Checked and rejected'.", False),
        ("The three price bands", True),
        (BAND_RULE + " The shortlist is grouped by band, Affordable first, and numbered within each band.", False),
        ("The result", True),
        (f"{len(SHORTLIST)} properties pass: {counts['Affordable']} affordable, {counts['Mid-range']} mid-range and "
         f"{counts['Premium']} premium. {len(CONFIRM_FIRST)} more have disputed room counts and are on 'Confirm first'. "
         f"{checked} names were checked in total; every one that failed is on 'Checked and rejected' with the reason. "
         f"A published figure covers a 650-guest reception at {yes['Affordable']} affordable and {yes['Mid-range']} "
         f"mid-range properties.", False),
        ("What was searched", True),
        ("Each hotel's own website, brochure or factsheet first. Then wedding-venue directories (WedMeGood, Weddingz, "
         "WeddingWire, VenueBookingz, The Wedding Company, Cvent, Spalba, venue.events) for lawn and hall sizes, and "
         "TripAdvisor and Google listings for ratings. Where a figure comes from a directory and not the hotel, the "
         "row says so.", False),
        ("How to read the shortlist", True),
        ("'Seats 650' is green where a published figure covers it, amber where the space is big enough but the hotel has "
         "not printed a seated number, and red where the published figure is too small. 'Room rate seen online' is "
         "what booking sites showed for a standard room, with the site named; where a site publishes a dearest "
         "month, it is quoted, because February is peak season. 'Plate or package price signal' is what a directory "
         "listed. Neither is a quote. Yellow columns are for the call and are empty.", False),
        ("Call first: affordable", True),
        ("jüSTa Rajputana (54 rooms, 4-star, rooms ₹3,500–5,200), Padmini Bagh (57, 4-star, a 13,000 sq ft hall), "
         "Rupis (70, 3-star, publishes 3- and 4-sharing rooms), Labh Garh (80, 4-star), Oriental Palace (48, 3-star, "
         "64,000 sq ft of lawns), jüSTa Ssatva (50, 4-star), The Fern Residency (55, 4-star, TA 4.7). Each publishes "
         "a lawn big enough for 650 by guest count or by area. Padmini Bagh's own pages contradict each other on its "
         "lawn, so ask which lawn and what seated layout.", False),
        ("Call first: mid-range", True),
        ("Amargarh (53 rooms, 4-star but priced higher, a hall for 800) and The Royal Retreat (100+, 5-star, a hall "
         "for 550 at round tables). Both have the most credible indoor rain plans on the list.", False),
        ("Still not known for any of them", True),
        ("Whether 31 January to 2 February 2027 is free, and the wedding price. No hotel publishes either. "
         "Use the 'Ask every hotel' tab and write answers on 'Quote log'.", False),
    ]
    row = 4
    for text, bold in lines:
        cell = ws.cell(row, 1, text)
        cell.font = BOLD if bold else BODY
        cell.alignment = WRAP
        ws.row_dimensions[row].height = 20 if bold else 16 * (len(text) // 105 + 1)
        row += 1 if not bold else 1
        if not bold:
            row += 1


def build_shortlist(wb):
    ws = wb.create_sheet("Shortlist")
    ws.sheet_properties.tabColor = "C9A227"
    header_row(ws, [c[0] for c in COLUMNS])
    for i, item in enumerate(SHORTLIST, 1):
        r = i + 1
        base = IVORY if r % 2 == 0 else CREAM
        for c, (_, key, _) in enumerate(COLUMNS, 1):
            if key is None:
                cell = ws.cell(r, c, None)
                cell.fill = fill(YELLOW)
            else:
                cell = ws.cell(r, c, item.get(key, ""))
                cell.fill = fill(base)
            cell.font = BOLD if key in ("name", "band") else BODY
            cell.alignment = WRAP
            cell.border = BORDER
        col_650 = next(i for i, c in enumerate(COLUMNS, 1) if c[1] == "seats_650")
        ws.cell(r, col_650).fill = fill(fit_colour(item["seats_650"]))
        for c in (1, 2):
            ws.cell(r, c).fill = fill(BAND_FILL[item["band"]])
        ws.row_dimensions[r].height = 120
    for c, (_, _, w) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "D2"
    last = 1 + len(SHORTLIST)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"
    first_yellow = next(i for i, c in enumerate(COLUMNS, 1) if c[1] is None)
    dv = DataValidation(type="list", formula1='"Yes,No,On hold,Not asked"', allow_blank=True)
    dv.add(f"{get_column_letter(first_yellow)}2:{get_column_letter(first_yellow + 2)}{last}")
    ws.add_data_validation(dv)
    ws.cell(1, first_yellow).comment = Comment("Filled after the call. No hotel publishes this.", "sheet")


def build_rejected(wb):
    ws = wb.create_sheet("Checked and rejected")
    ws.sheet_properties.tabColor = MUTED
    heads = ["Property", "Area", "What was found", "Why it is not on the shortlist", "Source"]
    header_row(ws, heads)
    r = 2
    for name, area, found, reason, src in REJECTED:
        for c, v in enumerate((name, area, found, reason, src), 1):
            cell = ws.cell(r, c, v)
            cell.font = BOLD if c == 1 else BODY
            cell.alignment = WRAP
            cell.border = BORDER
            cell.fill = fill(IVORY if r % 2 == 0 else CREAM)
        ws.row_dimensions[r].height = 34
        r += 1
    note = ("Names only, with no address, room count or venue data found: " + ", ".join(NAME_ONLY)
            + ". These were on the earlier padded tabs and should not have been.")
    cell = ws.cell(r + 1, 1, note)
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=5)
    cell.font = NOTE
    cell.alignment = WRAP
    ws.row_dimensions[r + 1].height = 60
    for c, w in enumerate((40, 26, 44, 46, 36), 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:E{r - 1}"


def build_confirm(wb):
    ws = wb.create_sheet("Confirm first")
    ws.sheet_properties.tabColor = "F0A72A"
    heads = ["Property", "Area", "Phone", "Room count, by source", "Garden", "Hall", "Why it is not on the shortlist yet"]
    header_row(ws, heads)
    for r, row in enumerate(CONFIRM_FIRST, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.font = BOLD if c == 1 else BODY
            cell.alignment = WRAP
            cell.border = BORDER
            cell.fill = fill(IVORY if r % 2 == 0 else CREAM)
        ws.row_dimensions[r].height = 48
    for c, w in enumerate((30, 28, 24, 34, 34, 30, 36), 1):
        ws.column_dimensions[get_column_letter(c)].width = w


def build_quote_log(wb):
    ws = wb.create_sheet("Quote log")
    ws.sheet_properties.tabColor = "F0A72A"
    heads = [
        "Resort", "Called on", "Spoke to", "31 Jan 2027", "1 Feb 2027", "2 Feb 2027",
        "Rooms they will contract", "Rooms with two doubles", "06:00 check-in in writing?",
        "Lawn for the reception", "Seated dinner on that lawn", "Indoor seats for 180",
        "Indoor seats for 650", "Veg plate, Tuesday, ₹", "Minimum covers", "Lawn rental, ₹",
        "Room rate with breakfast, ₹", "Complimentary rooms", "Outside décor / DJ / alcohol",
        "Music cut-off", "GST included?", "Their total, ₹", "Advance and forfeit", "Follow-up",
    ]
    header_row(ws, heads)
    for i, item in enumerate(SHORTLIST, 2):
        for c in range(1, len(heads) + 1):
            cell = ws.cell(i, c, item["name"] if c == 1 else None)
            cell.font = BOLD if c == 1 else BODY
            cell.alignment = WRAP
            cell.border = BORDER
            cell.fill = fill(BAND_FILL[item["band"]] if c == 1 else YELLOW)
        ws.row_dimensions[i].height = 30
    ws.column_dimensions["A"].width = 40
    for c in range(2, len(heads) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 16
    ws.freeze_panes = "B2"
    last = 1 + len(SHORTLIST)
    dv = DataValidation(type="list", formula1='"Yes,No,On hold,Not asked"', allow_blank=True)
    dv.add(f"D2:F{last}")
    dv.add(f"I2:I{last}")
    dv.add(f"U2:U{last}")
    ws.add_data_validation(dv)


def build_questions(wb):
    ws = wb.create_sheet("Ask every hotel")
    ws.sheet_properties.tabColor = "6E1526"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Ask every hotel the same questions"
    ws["A1"].font = TITLE
    ws["A2"] = "Write the answers on Quote log. If it is not in an email, it is not the price."
    ws["A2"].font = NOTE
    for i, q in enumerate(QUESTIONS, 1):
        ws.cell(i + 3, 1, i).font = BOLD
        cell = ws.cell(i + 3, 2, q)
        cell.font = BODY
        cell.alignment = WRAP
        ws.row_dimensions[i + 3].height = 32
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 120


def main():
    wb = Workbook()
    build_start(wb)
    build_shortlist(wb)
    build_confirm(wb)
    build_rejected(wb)
    build_quote_log(wb)
    build_questions(wb)
    wb.properties.title = "Udaipur resorts, 1–2 February 2027"
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
